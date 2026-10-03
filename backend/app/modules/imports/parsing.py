"""Reading uploaded spreadsheets and turning messy cell values into clean ones."""

import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from app.modules.imports.models import DateOrder

MAX_ROWS = 10_000


class ParseError(ValueError):
    """A cell or file couldn't be understood. The message is shown to the user."""


Row = dict[str, object]


def read_table(data: bytes, filename: str) -> tuple[list[str], list[Row]]:
    """Returns (headers, rows). Rows are dicts keyed by header. Blank rows are dropped."""
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
        headers, raw = _read_xlsx(data)
    elif name.endswith(".xls"):
        raise ParseError("Old .xls files aren't supported. Save it as .xlsx or CSV and try again.")
    else:
        headers, raw = _read_csv(data)

    headers = _unique_headers(headers)
    if not any(headers):
        raise ParseError("The first row should contain column names.")
    rows = []
    for values in raw:
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
            continue
        rows.append(
            {h: (values[i] if i < len(values) else None) for i, h in enumerate(headers) if h}
        )
        if len(rows) > MAX_ROWS:
            raise ParseError(
                f"Files can have up to {MAX_ROWS:,} rows. Split it into smaller files."
            )
    if not rows:
        raise ParseError("The file has no data rows.")
    return [h for h in headers if h], rows


def _unique_headers(headers: list) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for h in headers:
        h = str(h).strip() if h is not None else ""
        if h:
            seen[h] = seen.get(h, 0) + 1
            if seen[h] > 1:
                h = f"{h} ({seen[h]})"
        out.append(h)
    return out


def _read_csv(data: bytes) -> tuple[list, list[list]]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - cp1252 decodes almost anything
        raise ParseError("Couldn't read the file's text encoding. Save it as UTF-8 CSV.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    all_rows = list(reader)
    if not all_rows:
        raise ParseError("The file is empty.")
    return all_rows[0], all_rows[1:]


def _read_xlsx(data: bytes) -> tuple[list, list[list]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for bad files
        raise ParseError("Couldn't open the Excel file. Is it a valid .xlsx?") from exc
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    if not rows:
        raise ParseError("The first sheet is empty.")
    return rows[0], rows[1:]


# --- Cell values ---------------------------------------------------------------


def text(v: object) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)  # Excel stores 9876543210 as 9876543210.0
    s = str(v).strip()
    return s or None


_MONTHS = {
    m: i
    for i, m in enumerate(
        ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"],
        start=1,
    )
}


def parse_date(v: object, order: DateOrder) -> date | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    iso = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T].*)?", s)  # ISO always wins
    if iso:
        return _make(int(iso.group(1)), int(iso.group(2)), int(iso.group(3)), s)
    # 5 Oct 2026 / 05-Oct-26 / October 5, 2026
    m = re.fullmatch(r"(\d{1,2})[\s\-/.]*([A-Za-z]{3,9})[\s\-/.,]*(\d{2,4})", s)
    if m and m.group(2)[:3].lower() in _MONTHS:
        return _make(_year(m.group(3)), _MONTHS[m.group(2)[:3].lower()], int(m.group(1)), s)
    m = re.fullmatch(r"([A-Za-z]{3,9})[\s\-/.]*(\d{1,2}),?[\s\-/.]*(\d{2,4})", s)
    if m and m.group(1)[:3].lower() in _MONTHS:
        return _make(_year(m.group(3)), _MONTHS[m.group(1)[:3].lower()], int(m.group(2)), s)
    m = re.fullmatch(r"(\d{1,4})[/\-.](\d{1,2})[/\-.](\d{1,4})", s)
    if m:
        a, b, c = (int(x) for x in m.groups())
        if len(m.group(1)) == 4:
            return _make(a, b, c, s)
        if order == DateOrder.MDY:
            return _make(_year(m.group(3)), a, b, s)
        return _make(_year(m.group(3)), b, a, s)  # DMY (and YMD falls back to DMY for d/m/y)
    raise ParseError(f'"{s}" isn\'t a date I understand')


def _year(s: str) -> int:
    y = int(s)
    return y + 2000 if y < 100 else y


def _make(y: int, m: int, d: int, original: str) -> date:
    try:
        return date(y, m, d)
    except ValueError as exc:
        raise ParseError(f'"{original}" isn\'t a valid date') from exc


def parse_decimal(v: object) -> Decimal | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    s = re.sub(r"[^\d.\-]", "", str(v).replace(",", ""))  # "₹1,800.00" -> "1800.00"
    if not s:
        if str(v).strip():
            raise ParseError(f'"{v}" isn\'t a number')
        return None
    try:
        return Decimal(s)
    except InvalidOperation as exc:
        raise ParseError(f'"{v}" isn\'t a number') from exc


def parse_float(v: object) -> float | None:
    d = parse_decimal(v)
    return float(d) if d is not None else None


def parse_choice(v: object, synonyms: dict[str, str], label: str) -> str | None:
    """Maps free text onto an allowed value using a synonyms table (keys lowercase)."""
    s = text(v)
    if s is None:
        return None
    key = re.sub(r"[\s_\-]+", " ", s.lower()).strip()
    if key in synonyms:
        return synonyms[key]
    for k, value in synonyms.items():  # "Weight loss & toning" -> weight_loss
        if len(k) > 3 and k in key:
            return value
    raise ParseError(f'"{s}" isn\'t a known {label}')


def parse_duration(v: object) -> tuple[int, str] | None:
    """'3 months', 'Quarterly', '1 yr', '30 days' -> (3, 'month')."""
    s = text(v)
    if s is None:
        return None
    key = s.lower().strip()
    named = {
        "monthly": (1, "month"),
        "quarterly": (3, "month"),
        "half yearly": (6, "month"),
        "half-yearly": (6, "month"),
        "semi annual": (6, "month"),
        "yearly": (1, "year"),
        "annual": (1, "year"),
        "weekly": (1, "week"),
    }
    if key in named:
        return named[key]
    m = re.fullmatch(r"(\d+)\s*([a-z]+)", key)
    if m:
        unit = m.group(2)
        for prefix, canonical in (("d", "day"), ("w", "week"), ("m", "month"), ("y", "year")):
            if unit.startswith(prefix):
                return int(m.group(1)), canonical
    raise ParseError(f'"{s}" isn\'t a duration like "3 months"')
